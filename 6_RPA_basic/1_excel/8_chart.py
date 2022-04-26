from openpyxl import load_workbook
wb=load_workbook("sample.xlsx")
ws=wb.active

from openpyxl.chart import BarChart, Reference, LineChart

# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3) #B2부터 C11까지의 데이터를 설정
# bar_chart = BarChart() #차트 종류 설정 (Bar, Line, Pie, ..)
# bar_chart.add_data(bar_value) #차트 데이터 추가

# ws.add_chart(bar_chart, "E1") #E1 위치에 차트 추가

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) #계열을 설정함(각 데이터의 첫번째값에서)
line_chart.title = "성적표" #차트 제목
line_chart.style = 20 #미리 정의된 스타일을 적용, 사용자가 개별 지정도 가능
line_chart.y_axis.title = "점수" #Y축의 제목
line_chart.x_axis.title = "번호" #X축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")

#구글에서 openpyxl 검색후 사이트 진입해서 더 자세한 공부 가능